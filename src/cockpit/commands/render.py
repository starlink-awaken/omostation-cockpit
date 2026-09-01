"""cockpit.commands.render — one-shot export of drafts to DOCX/PPTX/SVG (BET-Y1Q4-T8-03).

docx  — GB/T 9704-2012 red-header official document (renderers/gov_docx)
pptx  — 16:9 executive tech deck (dark-business / minimal-tech templates)
svg   — vector architecture diagram from ```diagram code blocks
test_export_formats — offline verify contract over a synthetic document
"""

from __future__ import annotations

import argparse
from pathlib import Path

from rich.console import Console
from rich.panel import Panel

console = Console()

SCHEMA = "cockpit.render.v1"
DEFAULT_TEMPLATES = {"docx": "standard-gov", "pptx": "dark-business", "svg": "boxes-arrows"}

# 16:9 deck palettes (done_when: 深色商务 / 极简科技 双模板)
PPTX_TEMPLATES = {
    "dark-business": {"bg": "0F1B2D", "fg": "F5F7FA", "accent": "D4AF37", "title": "微软雅黑"},
    "minimal-tech": {"bg": "FFFFFF", "fg": "1F2937", "accent": "2563EB", "title": "微软雅黑"},
}


def cmd_render_docx(args: argparse.Namespace) -> int:
    """Render Markdown to GB/T 9704-2012 DOCX (degrades to clean .md on template failure)."""
    from cockpit.renderers.gov_docx import degrade_markdown, parse_markdown, render_docx

    md_path = Path(getattr(args, "input", ""))
    template = getattr(args, "template", "") or DEFAULT_TEMPLATES["docx"]
    output = getattr(args, "output", "") or str(md_path.with_suffix(".docx"))
    if not md_path.is_file():
        console.print(f"[red]输入文件不存在: {md_path}[/red]")
        return 1
    md = md_path.read_text(encoding="utf-8")
    try:
        if template != "standard-gov":
            raise ValueError(f"unknown docx template: {template}")
        out = render_docx(parse_markdown(md), output)
        console.print(Panel(f"[green]GB/T 9704-2012 DOCX 已生成[/green]\n{out}", title="📄 Render DOCX"))
        return 0
    except Exception as exc:  # circuit_breaker: 模板/解析异常 → 干净 Markdown 降级
        out = degrade_markdown(md, str(Path(output).with_suffix(".fallback.md")))
        console.print(f"[yellow]模板解析异常 ({exc}) → 降级输出标准 Markdown: {out}[/yellow]")
        return 0


def cmd_render_pptx(args: argparse.Namespace) -> int:
    """Render Markdown headings/lists into a 16:9 executive deck."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    md_path = Path(getattr(args, "input", ""))
    template = getattr(args, "template", "") or DEFAULT_TEMPLATES["pptx"]
    output = getattr(args, "output", "") or str(md_path.with_suffix(".pptx"))
    if not md_path.is_file():
        console.print(f"[red]输入文件不存在: {md_path}[/red]")
        return 1
    palette = PPTX_TEMPLATES.get(template)
    if palette is None:
        console.print(f"[red]未知 pptx 模板: {template} (可用: {', '.join(PPTX_TEMPLATES)})[/red]")
        return 1

    import re

    md = md_path.read_text(encoding="utf-8")
    title = next((ln[2:].strip() for ln in md.splitlines() if ln.startswith("# ")), md_path.stem)
    sections: list[tuple[str, list[str]]] = []
    current: tuple[str, list[str]] = ("", [])
    for line in md.splitlines():
        if line.startswith("## "):
            if current[0]:
                sections.append(current)
            current = (line[3:].strip(), [])
        elif bul := re.match(r"^\s*([-*+]|\d+\.)\s+(.+)$", line):
            current[1].append(bul.group(2).strip())
        elif current[0] and line.strip():
            current[1].append(line.strip())
    if current[0]:
        sections.append(current)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
    # default template ships type="screen4x3"; correct the label to 16:9
    sld_sz = prs.slides._sldIdLst.getparent().find(
        ".//{http://schemas.openxmlformats.org/presentationml/2006/main}sldSz"
    )
    if sld_sz is not None:
        sld_sz.set("type", "screen16x9")
    blank = prs.slide_layouts[6]
    bg_c, fg_c, ac_c = (RGBColor.from_string(palette[k]) for k in ("bg", "fg", "accent"))

    def _add_slide(text: str, size: int, accent_bar: bool = False) -> None:
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = bg_c
        box = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(2.8))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.color.rgb = fg_c
        run.font.name = palette["title"]
        if accent_bar:
            bar = slide.shapes.add_shape(1, Inches(5.9), Inches(4.1), Inches(1.5), Inches(0.08))
            bar.fill.solid()
            bar.fill.fore_color.rgb = ac_c
            bar.line.fill.background()

    _add_slide(title, 40)  # cover
    for heading, bullets in sections:
        _add_slide(heading, 32, accent_bar=True)
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = bg_c
        box = slide.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(10.9), Inches(5.4))
        tf = box.text_frame
        tf.word_wrap = True
        for idx, item in enumerate(bullets):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = f"•  {item}"
            run.font.size = Pt(20)
            run.font.color.rgb = fg_c
            p.space_after = Pt(14)

    out = Path(output)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    console.print(Panel(f"[green]16:9 PPTX 已生成[/green]\n{out} (template: {template})", title="📊 Render PPTX"))
    return 0


def cmd_render_svg(args: argparse.Namespace) -> int:
    """Render ```diagram code blocks to vector box-arrow SVG."""
    from cockpit.renderers.gov_docx import parse_markdown

    md_path = Path(getattr(args, "input", ""))
    output = getattr(args, "output", "") or str(md_path.with_suffix(".svg"))
    if not md_path.is_file():
        console.print(f"[red]输入文件不存在: {md_path}[/red]")
        return 1
    model = parse_markdown(md_path.read_text(encoding="utf-8"))
    if not model.diagrams:
        console.print("[yellow]未找到 ```diagram 代码块，无 SVG 生成[/yellow]")
        return 0

    svgs: list[str] = []
    for block in model.diagrams:
        nodes, edges = [], []
        for line in block.splitlines():
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            if "->" in line:
                a, _, b = line.partition("->")
                edges.append((a.strip().strip('"'), b.strip().strip('"')))
            else:
                nodes.append(line.strip().strip('"'))
        for a, b in edges:  # edge endpoints are nodes too (diagram blocks may omit standalone lines)
            for endpoint in (a, b):
                if endpoint not in nodes:
                    nodes.append(endpoint)
        node_w, node_h, gap = 220, 64, 48
        positions = {n: (40 + i * (node_w + gap), 60) for i, n in enumerate(nodes)}
        width = 80 + len(nodes) * (node_w + gap) if nodes else 640
        height = 220 + (24 * len(edges) if edges else 0)
        parts = [
            f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" font-family="sans-serif">',
        ]
        for name, (x, y) in positions.items():
            label = name if len(name) <= 16 else name[:15] + "…"
            parts.append(
                f'<rect x="{x}" y="{y}" width="{node_w}" height="{node_h}" rx="10" '
                f'fill="#EEF2FF" stroke="#2563EB" stroke-width="2"/>'
                f'<text x="{x + node_w / 2}" y="{y + node_h / 2 + 6}" text-anchor="middle" '
                f'font-size="16" fill="#1F2937">{label}</text>'
            )
        for i, (a, b) in enumerate(edges):
            if a in positions and b in positions:
                x1 = positions[a][0] + node_w
                x2 = positions[b][0]
                y = positions[a][1] + node_h // 2
                yy = y + 40 + i * 24 if x2 < x1 else y
                parts.append(
                    f'<path d="M {x1} {y} C {(x1 + x2) / 2} {yy}, {(x1 + x2) / 2} {yy}, {x2} {yy}" '
                    f'fill="none" stroke="#6B7280" stroke-width="2" marker-end="url(#arrow)"/>'
                )
        parts.append(
            '<defs><marker id="arrow" markerWidth="10" markerHeight="8" refX="9" refY="4" orient="auto">'
            '<path d="M0,0 L10,4 L0,8 z" fill="#6B7280"/></marker></defs></svg>'
        )
        svgs.append("\n".join(parts))

    out = Path(output)
    out.parent.mkdir(parents=True, exist_ok=True)
    if len(svgs) == 1:
        out.write_text(svgs[0], encoding="utf-8")
    else:  # multi-diagram: number the outputs
        out = out.with_suffix("")
        for i, svg in enumerate(svgs):
            (Path(f"{out}-{i + 1}.svg")).write_text(svg, encoding="utf-8")
    console.print(Panel(f"[green]矢量架构图 SVG 已生成[/green]\n{output}", title="🗂 Render SVG"))
    return 0


def _synthetic_doc() -> str:
    """Synthetic red-header document exercising every renderer branch."""
    return """# 关于推进数字医疗健康服务试点工作的通知

> 国卫办发布〔2026〕18号
> 2026年9月1日

各省、自治区、直辖市卫生健康委：

为深入推进健康中国战略，现就数字医疗健康服务试点工作通知如下。

## 一、试点安排

- 北京市率先开展医疗大模型临床辅助试点
- 上海市推进医疗数据要素互联互通示范区建设
- 广东省建设区域医疗健康大数据平台

## 二、保障措施

各试点地区要压实责任，确保数据安全与应用规范。

```diagram
卫健委 -> 试点医院
试点医院 -> 区域平台
区域平台 -> 卫健委
```
"""


def cmd_render_test(args: argparse.Namespace) -> int:
    """Offline verify contract: synthetic doc → DOCX/PPTX/SVG, assert GB/T key params."""
    import json
    import tempfile
    import zipfile

    from cockpit.renderers.gov_docx import GOV_SPEC, parse_markdown

    md = _synthetic_doc()
    tmp = Path(tempfile.mkdtemp(prefix="render-test-"))
    md_path = tmp / "doc.md"
    md_path.write_text(md, encoding="utf-8")

    ns = argparse.Namespace(input=str(md_path), output=str(tmp / "doc.docx"), template="standard-gov")
    rc_docx = cmd_render_docx(ns)
    ns = argparse.Namespace(input=str(md_path), output=str(tmp / "deck.pptx"), template="dark-business")
    rc_pptx = cmd_render_pptx(ns)
    ns = argparse.Namespace(input=str(md_path), output=str(tmp / "diag.svg"))
    rc_svg = cmd_render_svg(ns)

    checks = {"docx_generated": rc_docx == 0 and (tmp / "doc.docx").exists()}
    if checks["docx_generated"]:
        with zipfile.ZipFile(tmp / "doc.docx") as z:
            document = z.read("word/document.xml").decode("utf-8")
        m = GOV_SPEC["margin_cm"]
        # python-docx rounds Cm→twips half-up; mirror the same conversion
        checks["gb_margin_top"] = str(round(m["top"] * 566.929)) in document
        checks["gb_title_font"] = GOV_SPEC["title_font"]["east_asia"] in document
        checks["gb_body_font"] = GOV_SPEC["body_font"]["east_asia"] in document
        checks["gb_line_pitch_exact"] = 'w:lineRule="exact"' in document
    checks["pptx_generated"] = rc_pptx == 0 and (tmp / "deck.pptx").exists()
    if checks["pptx_generated"]:
        with zipfile.ZipFile(tmp / "deck.pptx") as z:
            slide1 = z.read("ppt/slides/slide1.xml").decode("utf-8")
            presentation = z.read("ppt/presentation.xml").decode("utf-8")
        # 16:9 ≈ 12192000 EMU wide (python-pptx Inches(13.333) = 12191695)
        checks["pptx_16_9"] = "screen4x3" not in presentation and "12191" in presentation
        checks["pptx_dark_bg"] = "0F1B2D" in slide1
    checks["svg_generated"] = rc_svg == 0 and (tmp / "diag.svg").exists()
    if checks["svg_generated"]:
        checks["svg_vector"] = "<svg" in (tmp / "diag.svg").read_text(encoding="utf-8")

    fidelity = round(sum(checks.values()) / len(checks), 3)
    report = {"schema": SCHEMA, "checks": checks, "fidelity": fidelity, "outdir": str(tmp)}
    print(json.dumps(report, ensure_ascii=False, indent=1))
    return 0 if fidelity >= 0.95 else 1


def cmd_render(args: argparse.Namespace) -> int:
    """Dispatch render subcommand."""
    sub = getattr(args, "render_command", None)
    dispatch = {
        "docx": cmd_render_docx,
        "pptx": cmd_render_pptx,
        "svg": cmd_render_svg,
        "test_export_formats": cmd_render_test,
    }
    if sub in dispatch:
        return dispatch[sub](args)
    console.print("[red]未知 render 子命令[/red] (可用: docx / pptx / svg / test_export_formats)")
    return 1
